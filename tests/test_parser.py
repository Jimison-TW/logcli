import logging
import subprocess
from collections import Counter

import pytest
from openpyxl import load_workbook

from logcli.logging_config import setup_logging  # import 加這行
from logcli.models import LogStats  # import 要加
from logcli.parser import (
    LogParseError,
    count_levels,
    extract_fields,
    extract_level,
    main,
    parse_line,
    read_lines,
    show_top_n,
    write_csv_report,
    write_excel_report,
)


@pytest.mark.parametrize(
    "log, expected",
    [
        ("2026-07-22T10:00:00 ERROR 資料庫 連線 失敗", "ERROR"),
        ("2026-07-22T10:00:02 WARN 記憶體偏高", "WARN"),
        ("2026-07-22T10:00:01 INFO 服務啟動", "INFO"),
        ("2026-07-22T10:00:00 CRITICAL 等級不存在", None),
        ("2026-07-22T10:00:00 INFO 等級不存在ERROR", "ERROR"),
    ],
)
def test_extract(log, expected):
    assert extract_level(log) == expected


@pytest.mark.parametrize(
    "log_list, expected",
    [
        ([], Counter([])),
        (["2026-07-22T10:00:00 CRITICAL 等級不存在"], Counter([])),
        (
            ["2026-07-22T10:00:01 INFO 服務啟動", "2026-07-22T10:00:00 INFO 等級不存在"],
            Counter({"INFO": 2}),
        ),
        (["2026-07-22T10:00:00 INFO 等級不存在ERROR"], Counter({"ERROR": 1})),
    ],
)
def test_counter(log_list, expected):
    assert count_levels(log_list) == expected


def test_counter_missing_level():
    result = count_levels(
        ["2026-07-22T10:00:01 INFO 服務啟動", "2026-07-22T10:00:00 INFO 等級不存在"]
    )
    assert result["ERROR"] == 0


def test_read_lines(sample_log_file):
    result = read_lines(sample_log_file)
    assert len(result) == 2


def test_read_lines_is_dir(tmp_path):
    with pytest.raises(FileExistsError):
        read_lines(tmp_path)  # tmp_path 本身是「目錄」不是檔 → 觸發那條分支


def test_read_lines_missing(tmp_path):
    missing = tmp_path / "nope.log"
    with pytest.raises(FileNotFoundError):
        read_lines(missing)


def test_read_lines_empty(tmp_path):
    empty = tmp_path / "empty.log"  # ← 在目錄底下指定一個檔名
    empty.write_text("", encoding="utf-8")  # ← 造出這個檔,內容你想想「空」是什麼
    result = read_lines(empty)  # ← 讀「檔」,不是讀「目錄」
    assert result == []


def test_write_csv(tmp_path):
    out = tmp_path / "out.csv"
    # rows = [
    #     "2026-07-22T10:00:00 ERROR 資料庫 連線 失敗",
    #     "2026-07-22T10:00:00 CRITICAL 等級不存在",
    # ]
    rows = [("ERROR", 1), ("CRITICAL", 1)]
    write_csv_report(rows, out)
    lines = out.read_text(encoding="utf-8").splitlines()
    assert lines[0] == "level,count"
    assert lines[1] == "ERROR,1"


def test_write_xlsx(tmp_path):
    out = tmp_path / "out.xlsx"
    # rows = [
    #     "2026-07-22T10:00:00 ERROR 資料庫 連線 失敗",
    #     "2026-07-22T10:00:00 CRITICAL 等級不存在",
    # ]
    rows = [("ERROR", 1), ("CRITICAL", 1)]
    write_excel_report(rows=rows, output=out)
    wb = load_workbook(out)
    ws = wb.active
    assert ws is not None
    all_rows = list(ws.iter_rows(values_only=True))
    assert all_rows[0] == ("level", "count")
    assert all_rows[1] == ("ERROR", 1)


def test_parse_line_normal():
    result = parse_line("2026-07-22T10:00:02 WARN 記憶體偏高")
    assert result.timestamp.isoformat() == "2026-07-22T10:00:02"
    assert result.level == "WARN"
    assert result.message == "記憶體偏高"


def test_parse_line_column_missing():
    with pytest.raises(LogParseError) as exc_info:
        parse_line("只有兩段 INFO")
    assert "訊息字串缺少錯誤內容" == exc_info.value.reason


def test_parse_line_error_time():
    with pytest.raises(LogParseError) as exc_info:
        parse_line("not-a-time ERROR 時間格式壞掉")
    assert "timestamp 格式錯誤" == exc_info.value.reason


def test_parse_line_wrong_level():
    with pytest.raises(LogParseError) as exc_info:
        parse_line("2026-07-22T10:00:00 CRITICAL 等級不存在")
    assert "level 不在 INFO/WARN/ERROR 之內" == exc_info.value.reason


def test_extract_fields_normal():
    result = extract_fields("2026-07-22T10:00:00 ERROR [192.168.0.1] code=E4001 資料庫連線失敗")
    assert result.ip == "192.168.0.1"
    assert result.code == "E4001"


def test_extract_fields_simple():
    result = extract_fields("2026-07-22T10:00:00 INFO 服務啟動")
    assert result.ip is None
    assert result.code is None


def test_extract_fields_wrong():
    with pytest.raises(LogParseError) as exc_info:
        extract_fields("garbage")
    assert "找不到符合的欄位" == exc_info.value.reason


def test_extract_fields_ilegal():
    with pytest.raises(LogParseError) as exc_info:
        extract_fields("2026-07-22T10:00:00 CRITICAL msg")
    assert "level 不合法" == exc_info.value.reason


def test_extract_fields_error_format():
    with pytest.raises(LogParseError) as exc_info:
        extract_fields("2026-13-45T00:00:00 INFO msg")
    assert "timestamp 格式錯誤" == exc_info.value.reason


def test_error_rate_zero_total():
    assert LogStats(total=0).error_rate == 0.0  # total=0 → 走 else 0.0,不會除以零


def test_show_top_n_plain(capsys):
    # rows = [
    #     "2026-07-22T10:00:00 ERROR 資料庫連線失敗",
    #     "2026-07-22T10:00:00 ERROR 磁碟寫入失敗",
    #     "2026-07-22T10:00:00 WARN 記憶體偏高",
    # ]
    counter = Counter({"ERROR": 2, "WARN": 1})
    show_top_n(counter, top_count=None, with_json=False)
    captured = capsys.readouterr()
    assert captured.out == "ERROR: 2\n" + "WARN: 1\n"  # ← 你來填


def test_show_top_n_table(capsys):
    # rows = [
    #     "2026-07-22T10:00:00 ERROR 資料庫連線失敗",
    #     "2026-07-22T10:00:00 ERROR 磁碟寫入失敗",
    #     "2026-07-22T10:00:00 WARN 記憶體偏高",
    # ]
    counter = Counter({"ERROR": 2, "WARN": 1})
    show_top_n(counter, top_count=None, with_json=False, fmt="table")
    captured = capsys.readouterr()
    assert "ERROR" in captured.out
    assert "level" in captured.out


def test_show_top_n_json(capsys):
    counter = Counter({"ERROR": 1})
    show_top_n(counter, top_count=None, with_json=True)
    captured = capsys.readouterr()
    assert captured.out == '[["ERROR", 1]]\n'


"""
logging.basicConfig() 有一條隱藏規則:如果 root logger「已經有 handler」,它就什麼都不做(no-op)。

而 pytest 為了攔截 log,啟動時早就往 root logger 塞了自己的 handler。所以你的 setup_logging 一呼叫 
basicConfig,它看到「已經有 handler 了」→ 直接擺爛、連 level 都沒設 → root 維持 WARNING → 斷言炸。

這是面試等級的坑:basicConfig 是「一次性、且只在乾淨狀態下生效」的。你在 production 沒踩到,是因為 CLI
啟動時 root 還是乾淨的;一進 pytest 環境就現形。
"""
# def test_setup_logging_verbose():
#     setup_logging(verbose=True)
#     assert logging.getLogger().level == logging.DEBUG  # verbose=True → root logger 設成 DEBUG


def test_setup_logging_verbose():
    root = logging.getLogger()
    original = root.handlers[:]  # ① 存檔:把 pytest 塞的 handler 備份起來
    root.handlers.clear()  # ② 清乾淨 → basicConfig 才不會被 no-op 擋掉
    try:
        setup_logging(verbose=True)
        assert root.level == logging.DEBUG  # 現在 basicConfig 真的生效了
    finally:
        root.handlers[:] = original  # ③ 還原:不管測試過不過都復原,不污染其他測試


@pytest.mark.smoke
def test_cli_end_to_end(make_test_case):
    # log = tmp_path / "test.log"
    # log.write_text(
    #     "2026-07-22T10:00:00 ERROR 資料庫連線失敗\n"
    #     "2026-07-22T10:00:01 INFO 服務啟動\n"
    #     "2026-07-22T10:00:02 CRITICAL 等級不存在",
    #     encoding="utf-8",
    # )
    # result = subprocess.run(["logcli", "parse", "--top", "5", log],
    # capture_output=True, text=True)
    # assert result.returncode == 0
    # assert "ERROR: 1" in result.stdout

    case = make_test_case(
        "cli_end_to_end",
        "2026-07-22T10:00:00 ERROR 資料庫連線失敗\n"
        "2026-07-22T10:00:01 INFO 服務啟動\n"
        "2026-07-22T10:00:02 CRITICAL 等級不存在",
    )
    result = subprocess.run(["logcli", "parse", "--top", "5", case], capture_output=True, text=True)
    assert result.returncode == 0
    assert "ERROR: 1" in result.stdout


@pytest.mark.smoke
def test_main_in_process(make_test_case, capsys):
    # log = tmp_path / "test.log"
    # log.write_text(
    #     "2026-07-22T10:00:00 ERROR 資料庫連線失敗\n"
    #     "2026-07-22T10:00:01 INFO 服務啟動\n"
    #     "2026-07-22T10:00:02 CRITICAL 等級不存在",
    #     encoding="utf-8",
    # )
    case = make_test_case(
        "main_in_process",
        "2026-07-22T10:00:00 ERROR 資料庫連線失敗\n"
        "2026-07-22T10:00:01 INFO 服務啟動\n"
        "2026-07-22T10:00:02 CRITICAL 等級不存在",
    )
    main(["parse", "--top", "5", str(case)])
    assert capsys.readouterr().out == "ERROR: 1\nINFO: 1\n"


def test_until_includes_that_whole_day(make_test_case, capsys):
    """--until 2026-08-14 的規格:含 08-14 整天、排 08-15 起。"""
    # log = tmp_path / "test.log"
    # log.write_text(
    #     "2026-08-14T23:59:00 ERROR 當天最後一刻\n"  # 界內:當天最後一秒 → 期望【留】
    #     "2026-08-15T00:00:00 WARN 隔天第一刻\n",  # 界外:隔天第一秒 → 期望【排】
    #     encoding="utf-8",
    # )
    case = make_test_case(
        "until_includes_that_whole_day",
        "2026-08-14T23:59:00 ERROR 當天最後一刻\n"  # 界內:當天最後一秒 → 期望【留】
        "2026-08-15T00:00:00 WARN 隔天第一刻\n",  # 界外:隔天第一秒 → 期望【排】
    )
    main(["parse", "--until", "2026-08-14", "--top", "5", str(case)])
    assert capsys.readouterr().out == "ERROR: 1\n"


def test_since_includes_that_whole_day(make_test_case, capsys):
    """--since 2026-08-14 的規格:含 08-14 整天、排 08-13 以前。"""
    # log = tmp_path / "test.log"
    # log.write_text(
    #     "2026-08-13T23:59:00 ERROR 當天最後一刻\n"  # 界外:前一天最後一分鐘 → 期望【排除】
    #     "2026-08-14T00:00:00 WARN 隔天第一刻\n",  # 界內:當天第一秒 → 期望【留】
    #     encoding="utf-8",
    # )
    case = make_test_case(
        "since_includes_that_whole_day",
        "2026-08-13T23:59:00 ERROR 當天最後一刻\n"  # 界外:前一天最後一分鐘 → 期望【排除】
        "2026-08-14T00:00:00 WARN 隔天第一刻\n",  # 界內:當天第一秒 → 期望【留】
    )
    main(["parse", "--since", "2026-08-14", "--top", "5", str(case)])
    assert capsys.readouterr().out == "WARN: 1\n"
